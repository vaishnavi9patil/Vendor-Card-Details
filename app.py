import csv
import io
import json
import sqlite3
import threading
import time
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

from flask import (
    Flask, Response, abort, flash, jsonify, redirect, render_template,
    request, send_from_directory, url_for,
)
from PIL import Image, ImageOps

import ai_interpreter
import scan_card
import sheet_sync
from ai_interpreter import AIUnavailable, interpret_vendor_text, local_extract
from dedupe import Matcher, merge_records
from importer import TARGETS, ImportFileError, guess_target, read_rows, row_to_vendor
from scan_card import extract_text, ocr  # OCR engine is created once, in scan_card

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
IMPORT_DIR = UPLOAD_DIR / "imports"
DB_PATH = BASE_DIR / "vendors.db"

UPLOAD_DIR.mkdir(exist_ok=True)
IMPORT_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = "change-this-local-development-key"
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB per upload

@app.template_filter("localtime")
def localtime(value):
    """SQLite stores UTC; show it in this computer's time zone."""
    if not value:
        return "—"
    try:
        utc = datetime.strptime(str(value)[:19], "%Y-%m-%d %H:%M:%S")
        local = utc.replace(tzinfo=timezone.utc).astimezone()
        return local.strftime("%d %b %Y, %H:%M")
    except ValueError:
        return value


LIST_FIELDS = (
    "phone_numbers", "emails", "destinations_covered", "websites",
    "other_contacts",
)
ALL_FIELDS = [
    "company", "person_name", "designation", "phone_numbers", "emails",
    "destinations_covered", "websites", "address", "services_notes",
    "other_contacts",
]


# ------------------------------------------------------------ database

def init_db():
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS vendors (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                company TEXT,
                person_name TEXT,
                designation TEXT,
                phone_numbers TEXT,
                emails TEXT,
                destinations_covered TEXT,
                websites TEXT,
                address TEXT,
                services_notes TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        # Columns added later - created automatically on older databases
        existing = {row[1] for row in conn.execute("PRAGMA table_info(vendors)")}
        if "other_contacts" not in existing:
            conn.execute("ALTER TABLE vendors ADD COLUMN other_contacts TEXT")
        if "updated_at" not in existing:
            conn.execute("ALTER TABLE vendors ADD COLUMN updated_at TIMESTAMP")
        if "ocr_text" not in existing:
            # Raw OCR text, so AI can be re-run later without re-scanning
            conn.execute("ALTER TABLE vendors ADD COLUMN ocr_text TEXT")
        if "ai_status" not in existing:
            # 'ok' = read by Gemini, 'pending' = basic extraction, retry AI
            conn.execute("ALTER TABLE vendors ADD COLUMN ai_status TEXT")
        if "reviewed" not in existing:
            # 1 = a person has checked / edited the record (ready for CRM)
            conn.execute("ALTER TABLE vendors ADD COLUMN reviewed INTEGER DEFAULT 0")
        if "front_image" not in existing:
            conn.execute("ALTER TABLE vendors ADD COLUMN front_image TEXT")
        if "back_image" not in existing:
            conn.execute("ALTER TABLE vendors ADD COLUMN back_image TEXT")
        if "source" not in existing:
            # 'scan', 'import' or 'manual'
            conn.execute("ALTER TABLE vendors ADD COLUMN source TEXT DEFAULT 'scan'")
        if "created_by" not in existing:
            conn.execute("ALTER TABLE vendors ADD COLUMN created_by TEXT")

        conn.execute(
            "CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT)"
        )
        # Which vendors are linked to a row in the Excel sheet (see sheet_sync.py)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS sheet_links (
                vendor_id INTEGER PRIMARY KEY,
                app_hash TEXT,
                sheet_hash TEXT,
                deleted INTEGER DEFAULT 0,
                removed INTEGER DEFAULT 0,
                synced_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)


# ------------------------------------------------------------ settings

DEFAULT_SETTINGS = {
    # Profile
    "profile_name": "",
    "profile_email": "",
    "profile_phone": "",
    "profile_designation": "",
    # Organisation
    "company_name": "Khanna Travels",
    # AI / scanning
    "gemini_model": "gemini-3.8-flash",
    "scan_quality": "fast",          # fast | balanced | accurate
    # Export
    "export_reviewed_only": "0",
    # Excel (Microsoft 365) sync
    "ms_client_id": "",
    "ms_tenant_id": "",
    "sheet_url": "https://khannatravels-my.sharepoint.com/:x:/r/personal/operations_khannatravels_com/_layouts/15/Doc.aspx?sourcedoc=%7BE15D2D55-8A50-4358-9835-C8BAE29172ED%7D&file=Card%20Details.xlsx&action=default&mobileredirect=true",
    "sheet_worksheet": "CARD_DETAILS",
    "sheet_ids": "",
    "sync_enabled": "1",
    "sync_hour": "10",
    "sync_last_at": "",
    "sync_last_ok": "",
    "sync_last_report": "",
    "sync_last_auto_date": "",
}

SCAN_QUALITY = {"fast": 960, "balanced": 1280, "accurate": 1600}


def get_settings():
    settings = dict(DEFAULT_SETTINGS)
    with sqlite3.connect(DB_PATH) as conn:
        for key, value in conn.execute("SELECT key, value FROM settings"):
            settings[key] = value
    return settings


def save_settings(values):
    with sqlite3.connect(DB_PATH) as conn:
        for key, value in values.items():
            if key in DEFAULT_SETTINGS:
                conn.execute(
                    "INSERT INTO settings (key, value) VALUES (?, ?) "
                    "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                    (key, value),
                )
    apply_settings()


def apply_settings():
    """Push saved settings into the AI and OCR modules."""
    settings = get_settings()
    model = settings["gemini_model"].strip()
    if model:
        ai_interpreter.MODELS[0] = model
    scan_card.DET_SIDE_FAST = SCAN_QUALITY.get(settings["scan_quality"], 960)


def current_user():
    return get_settings()["profile_name"].strip() or None


def _row_to_vendor(row):
    vendor = dict(row)
    for field in LIST_FIELDS:
        try:
            vendor[field] = json.loads(vendor.get(field) or "[]")
        except (ValueError, TypeError):
            vendor[field] = []
    if vendor.get("ai_status") == "pending":
        vendor["status"] = "pending"
    elif vendor.get("reviewed"):
        vendor["status"] = "reviewed"
    else:
        vendor["status"] = "new"
    return vendor


def get_vendor(vendor_id):
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT * FROM vendors WHERE id = ?", (vendor_id,)).fetchone()
    return _row_to_vendor(row) if row else None


def find_existing_vendor(conn, vendor):
    """Same person at the same company, or a shared email address."""
    conn.row_factory = sqlite3.Row

    if vendor.get("company") and vendor.get("person_name"):
        row = conn.execute(
            """SELECT * FROM vendors
               WHERE lower(trim(company)) = lower(trim(?))
                 AND lower(trim(person_name)) = lower(trim(?))
               ORDER BY id DESC LIMIT 1""",
            (vendor["company"], vendor["person_name"]),
        ).fetchone()
        if row:
            return _row_to_vendor(row)

    if vendor.get("company") and not vendor.get("person_name"):
        # Company without a contact name: same company, also without a name
        row = conn.execute(
            """SELECT * FROM vendors
               WHERE lower(trim(company)) = lower(trim(?))
                 AND trim(COALESCE(person_name, '')) = ''
               ORDER BY id DESC LIMIT 1""",
            (vendor["company"],),
        ).fetchone()
        if row:
            return _row_to_vendor(row)

    for email in vendor.get("emails", []):
        row = conn.execute(
            "SELECT * FROM vendors WHERE emails LIKE ? ORDER BY id DESC LIMIT 1",
            (f'%"{email}"%',),
        ).fetchone()
        if row:
            existing = _row_to_vendor(row)
            # Only a match if it is the same person (or no person on either)
            same_person = (
                not vendor.get("person_name")
                or not existing.get("person_name")
                or existing["person_name"].lower() == vendor["person_name"].lower()
            )
            if same_person:
                return existing
    return None


def merge_vendor(old, new):
    """The latest scan wins; fields it left empty keep their old value."""
    merged = {}
    for field in ALL_FIELDS:
        empty = [] if field in LIST_FIELDS else ""
        merged[field] = new.get(field) or old.get(field) or empty
    return merged


def _field_values(vendor):
    values = []
    for field in ALL_FIELDS:
        value = vendor.get(field, [] if field in LIST_FIELDS else "")
        if field in LIST_FIELDS:
            value = json.dumps(value, ensure_ascii=False)
        values.append(value)
    return values


def update_vendor(conn, vendor_id, vendor, ocr_text, ai_status):
    assignments = ", ".join(f"{f} = ?" for f in ALL_FIELDS)
    conn.execute(
        f"UPDATE vendors SET {assignments}, ocr_text = ?, ai_status = ?, "
        "updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        _field_values(vendor) + [ocr_text, ai_status, vendor_id],
    )


def save_vendor(vendor, ocr_text="", ai_status="ok", images=(None, None),
                source="scan", reviewed=False, on_duplicate="update"):
    """Insert a new vendor, or update the existing record for the same
    person. Returns (result, vendor_id); result is 'added', 'updated' or
    'unchanged'."""
    front_image, back_image = images
    with sqlite3.connect(DB_PATH) as conn:
        existing = find_existing_vendor(conn, vendor)

        if existing:
            if on_duplicate == "skip":
                return "unchanged", existing["id"]
            # A basic (no-AI) scan must not overwrite a good AI record
            if ai_status == "pending" and existing.get("ai_status") in (None, "ok"):
                return "unchanged", existing["id"]
            update_vendor(conn, existing["id"], merge_vendor(existing, vendor),
                          ocr_text or existing.get("ocr_text") or "", ai_status)
            conn.execute(
                "UPDATE vendors SET reviewed = ?, "
                "front_image = COALESCE(?, front_image), "
                "back_image = COALESCE(?, back_image) WHERE id = ?",
                (1 if reviewed else 0, front_image, back_image, existing["id"]),
            )
            return "updated", existing["id"]

        cursor = conn.execute(
            f"INSERT INTO vendors ({', '.join(ALL_FIELDS)}, ocr_text, ai_status, "
            "reviewed, front_image, back_image, source, created_by) "
            f"VALUES ({', '.join('?' for _ in ALL_FIELDS)}, ?, ?, ?, ?, ?, ?, ?)",
            _field_values(vendor) + [ocr_text, ai_status, 1 if reviewed else 0,
                                     front_image, back_image, source, current_user()],
        )
        return "added", cursor.lastrowid


def get_vendors(search="", status="", source=""):
    vendors = _query_vendors(search)
    if status:
        vendors = [v for v in vendors if v["status"] == status]
    if source:
        vendors = [v for v in vendors if (v.get("source") or "scan") == source]
    return vendors


def _query_vendors(search=""):
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        if search:
            pattern = f"%{search}%"
            where = " OR ".join(f"{f} LIKE ?" for f in ALL_FIELDS)
            rows = conn.execute(
                f"SELECT * FROM vendors WHERE {where} ORDER BY id DESC",
                [pattern] * len(ALL_FIELDS),
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM vendors ORDER BY id DESC").fetchall()
    return [_row_to_vendor(row) for row in rows]


# ------------------------------------------------------------- images

def prepare_image(file_storage):
    """Save an upright, resized copy of the upload and return its path."""
    destination = UPLOAD_DIR / f"{uuid.uuid4().hex}.jpg"
    with Image.open(file_storage.stream) as original:
        image = ImageOps.exif_transpose(original).convert("RGB")
        image.thumbnail((2000, 2000))
        image.save(destination, "JPEG", quality=92)
    return destination


# -------------------------------------------------------------- routes

@app.route("/")
def index():
    search = request.args.get("search", "").strip()
    status = request.args.get("status", "").strip()
    all_matching = get_vendors(search)
    vendors = [v for v in all_matching if not status or v["status"] == status]
    counts = {
        "all": len(all_matching),
        "new": sum(v["status"] == "new" for v in all_matching),
        "reviewed": sum(v["status"] == "reviewed" for v in all_matching),
        "pending": sum(v["status"] == "pending" for v in all_matching),
    }
    return render_template("index.html", vendors=vendors, search=search,
                           status=status, counts=counts)


MAX_CARDS_PER_SCAN = 5


def scan_one_card(front, back=None):
    """Read one business card (front + optional back) and save it.
    Returns a dict describing the result (used by the page and the API)."""
    start = time.perf_counter()

    sides = [("FRONT", prepare_image(front))]
    if back and back.filename:
        sides.append(("BACK", prepare_image(back)))
    images = (sides[0][1].name, sides[1][1].name if len(sides) > 1 else None)

    text_parts = []
    for label, path in sides:
        text = extract_text(path)
        if text.strip():
            if len(sides) > 1:
                text = f"--- {label} OF CARD ---\n{text}"
            text_parts.append(text)
    ocr_text = "\n\n".join(text_parts)

    ocr_seconds = time.perf_counter() - start
    print("\n--- RAW OCR TEXT ---")
    print(ocr_text)
    print(f"OCR time: {ocr_seconds:.1f}s")

    if not ocr_text:
        return {"ok": False, "category": "error",
                "message": "No text detected. Try a clearer, well-lit photo of the card."}

    ai_start = time.perf_counter()
    try:
        vendor = interpret_vendor_text(ocr_text)
        ai_status = "ok"
    except AIUnavailable as error:
        # Don't lose the scan: save what simple rules can find and
        # let the user re-run AI later without scanning again.
        print(f"AI unavailable, using basic extraction: {error}")
        vendor = local_extract(ocr_text)
        ai_status = "pending"
    ai_seconds = time.perf_counter() - ai_start
    print(f"AI interpretation time: {ai_seconds:.1f}s")
    print(json.dumps(vendor, indent=2, ensure_ascii=False))

    result, vendor_id = save_vendor(vendor, ocr_text, ai_status, images)
    name = vendor.get("company") or vendor.get("person_name") or "Vendor"
    timing = f"Read in {ocr_seconds:.0f}s, interpreted in {ai_seconds:.0f}s."
    if result == "unchanged":
        category, message = "info", f"{name} is already saved. Nothing was changed. {timing}"
        short = "Already saved - nothing changed"
    elif ai_status == "pending":
        category = "warning"
        message = (f"{name} saved with basic details only - the AI service is busy. "
                   f"Use 'Retry AI' in a few minutes. {timing}")
        short = "Saved with basic details (AI busy)"
    elif result == "updated":
        category = "success"
        message = f"{name} already existed - the record was updated. Please review the details. {timing}"
        short = "Already existed - record updated"
    else:
        category = "success"
        message = f"{name} saved. Please review the details below. {timing}"
        short = "Saved"

    return {
        "ok": True, "category": category, "message": message, "short": short,
        "result": result, "ai_status": ai_status, "vendor_id": vendor_id,
        "name": name, "person": vendor.get("person_name") or "",
        "url": url_for("vendor_detail", vendor_id=vendor_id),
        "seconds": round(time.perf_counter() - start),
    }


@app.route("/scan", methods=["GET", "POST"])
def scan():
    """Scan page. The page normally sends each card to /scan/card one by
    one; this POST is the simple fallback for a single card."""
    if request.method == "POST":
        front = request.files.get("front_1") or request.files.get("front")
        back = request.files.get("back_1") or request.files.get("back")
        if not front or not front.filename:
            flash("Please select the front image of the card.", "error")
            return redirect(url_for("scan"))
        try:
            outcome = scan_one_card(front, back)
        except Exception as error:
            app.logger.exception("Card scanning failed")
            flash(f"Scanning failed: {error}", "error")
            return redirect(url_for("scan"))
        flash(outcome["message"], outcome["category"])
        if outcome["ok"]:
            return redirect(outcome["url"])
        return redirect(url_for("scan"))

    return render_template("scan.html", max_cards=MAX_CARDS_PER_SCAN)


@app.route("/scan/card", methods=["POST"])
def scan_card_api():
    """Scan one card and answer with JSON (used for multi-card scanning)."""
    front = request.files.get("front")
    back = request.files.get("back")
    if not front or not front.filename:
        return jsonify({"ok": False, "category": "error",
                        "message": "The front image of the card is missing."}), 400
    try:
        return jsonify(scan_one_card(front, back))
    except Exception as error:
        app.logger.exception("Card scanning failed")
        return jsonify({"ok": False, "category": "error",
                        "message": f"Scanning failed: {error}"}), 500


@app.route("/scan/done", methods=["POST"])
def scan_done():
    """After a single-card scan the page asks us to show the usual message."""
    message = request.form.get("message", "")[:500]
    category = request.form.get("category", "success")
    if message:
        flash(message, category if category in ("success", "info", "warning", "error") else "info")
    target = request.form.get("next") or ""
    if not target.startswith("/") or target.startswith("//"):
        target = url_for("index")
    return redirect(target)


def _lines(name):
    """Textarea with one value per line -> list."""
    values = []
    for line in request.form.get(name, "").splitlines():
        line = " ".join(line.split())
        if line and line not in values:
            values.append(line)
    return values


@app.route("/vendor/<int:vendor_id>", methods=["GET", "POST"])
def vendor_detail(vendor_id):
    vendor = get_vendor(vendor_id)
    if not vendor:
        abort(404)

    if request.method == "POST":
        edited = {}
        for field in ALL_FIELDS:
            if field in LIST_FIELDS:
                edited[field] = _lines(field)
            else:
                edited[field] = " ".join(request.form.get(field, "").split())
        with sqlite3.connect(DB_PATH) as conn:
            # A person checked it, so it no longer waits for the AI
            ai_status = "manual" if vendor["ai_status"] == "pending" else (vendor["ai_status"] or "ok")
            update_vendor(conn, vendor_id, edited, vendor["ocr_text"], ai_status)
            conn.execute("UPDATE vendors SET reviewed = 1 WHERE id = ?", (vendor_id,))
        flash("Changes saved. The record is marked as reviewed.", "success")
        return redirect(url_for("vendor_detail", vendor_id=vendor_id))

    return render_template("vendor.html", vendor=vendor)


@app.route("/vendor/new", methods=["GET", "POST"])
def vendor_new():
    """Add a vendor by hand (no card)."""
    if request.method == "POST":
        vendor = {}
        for field in ALL_FIELDS:
            if field in LIST_FIELDS:
                vendor[field] = _lines(field)
            else:
                vendor[field] = " ".join(request.form.get(field, "").split())
        if not (vendor["company"] or vendor["person_name"]):
            flash("Please enter at least a company or a contact name.", "error")
            return render_template("vendor.html", vendor=vendor, is_new=True)
        result, vendor_id = save_vendor(vendor, ai_status="manual", source="manual",
                                        reviewed=True, on_duplicate="skip")
        if result == "unchanged":
            flash("This vendor already exists - opened the existing record.", "info")
        else:
            flash("Vendor added.", "success")
        return redirect(url_for("vendor_detail", vendor_id=vendor_id))

    empty = {f: ([] if f in LIST_FIELDS else "") for f in ALL_FIELDS}
    return render_template("vendor.html", vendor=empty, is_new=True)


@app.route("/vendor/<int:vendor_id>/delete", methods=["POST"])
def vendor_delete(vendor_id):
    vendor = get_vendor(vendor_id)
    if vendor:
        with sqlite3.connect(DB_PATH) as conn:
            conn.execute("DELETE FROM vendors WHERE id = ?", (vendor_id,))
            conn.execute("UPDATE sheet_links SET deleted = 1 WHERE vendor_id = ?", (vendor_id,))
        flash(f"{vendor['company'] or 'Vendor'} was deleted.", "info")
    return redirect(url_for("index"))


@app.route("/cards/<path:filename>")
def card_image(filename):
    return send_from_directory(UPLOAD_DIR, filename)


# ------------------------------------------------------ CRM export / API

CRM_COLUMNS = [
    "Record ID", "Company Name", "First Name", "Last Name", "Full Name",
    "Job Title", "Primary Phone", "Other Phones", "Primary Email",
    "Other Emails", "Website", "Address", "Destinations", "Notes",
    "Other Contacts", "Status", "Created", "Updated", "Lead Source",
]


def _iso_local(value):
    """UTC timestamp from SQLite -> local 'YYYY-MM-DD HH:MM' for the CRM."""
    if not value:
        return ""
    try:
        utc = datetime.strptime(str(value)[:19], "%Y-%m-%d %H:%M:%S")
        return utc.replace(tzinfo=timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return str(value)


SOURCE_LABELS = {
    "scan": "Business card scan",
    "import": "Excel import",
    "manual": "Added manually",
    "sheet": "Excel sheet",
}


def to_crm_row(v):
    """One vendor in the flat format most CRMs import (Zoho, HubSpot,
    Salesforce, Excel)."""
    parts = (v.get("person_name") or "").split()
    phones, emails, sites = v["phone_numbers"], v["emails"], v["websites"]
    return {
        "Record ID": v["id"],
        "Company Name": v.get("company") or "",
        "First Name": parts[0].title() if parts else "",
        "Last Name": " ".join(parts[1:]).title(),
        "Full Name": v.get("person_name") or "",
        "Job Title": v.get("designation") or "",
        "Primary Phone": phones[0] if phones else "",
        "Other Phones": "; ".join(phones[1:]),
        "Primary Email": emails[0] if emails else "",
        "Other Emails": "; ".join(emails[1:]),
        "Website": sites[0] if sites else "",
        "Address": v.get("address") or "",
        "Destinations": "; ".join(v["destinations_covered"]),
        "Notes": v.get("services_notes") or "",
        "Other Contacts": " | ".join(v["other_contacts"]),
        "Status": {"reviewed": "Reviewed", "new": "Not reviewed",
                   "pending": "AI pending"}.get(v["status"], ""),
        "Created": _iso_local(v.get("created_at")),
        "Updated": _iso_local(v.get("updated_at")),
        "Lead Source": SOURCE_LABELS.get(v.get("source") or "scan", "Business card scan"),
    }


def _export_selection():
    """Vendors chosen by ?search=, ?status=, ?reviewed=1 and ?source=."""
    status = request.args.get("status", "").strip()
    if request.args.get("reviewed") == "1":
        status = "reviewed"
    return get_vendors(
        request.args.get("search", "").strip(),
        status=status,
        source=request.args.get("source", "").strip(),
    )


@app.route("/export")
def export_page():
    settings = get_settings()
    all_vendors = get_vendors()
    counts = {
        "all": len(all_vendors),
        "reviewed": sum(v["status"] == "reviewed" for v in all_vendors),
        "new": sum(v["status"] == "new" for v in all_vendors),
    }
    return render_template("export.html", counts=counts, settings=settings)


@app.route("/export.csv")
def export_csv():
    """CSV for CRM import / Excel. Respects the current search and the
    'reviewed only' option."""
    vendors = _export_selection()

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=CRM_COLUMNS)
    writer.writeheader()
    for v in vendors:
        writer.writerow(to_crm_row(v))

    filename = f"vendors_{datetime.now():%Y-%m-%d}.csv"
    return Response(
        "\ufeff" + buffer.getvalue(),           # BOM so Excel shows UTF-8 correctly
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.route("/export.xlsx")
def export_xlsx():
    """Formatted Excel file (same columns as the CSV). Respects the current
    search and the 'reviewed only' option."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    vendors = _export_selection()

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"

    # Lists one per line inside the cell are easier to read in Excel
    multi_line = {"Other Phones", "Other Emails", "Destinations", "Other Contacts"}
    ws.append(CRM_COLUMNS)
    for v in vendors:
        row = to_crm_row(v)
        values = []
        for col in CRM_COLUMNS:
            value = row[col]
            if col in multi_line and isinstance(value, str):
                value = value.replace("; ", "\n").replace(" | ", "\n")
            values.append(value)
        ws.append(values)

    header_fill = PatternFill("solid", fgColor="09488E")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

    widths = {
        "Record ID": 10, "Company Name": 30, "First Name": 14, "Last Name": 18,
        "Full Name": 24, "Job Title": 28, "Primary Phone": 18, "Other Phones": 20,
        "Primary Email": 30, "Other Emails": 30, "Website": 26, "Address": 45,
        "Destinations": 22, "Notes": 45, "Other Contacts": 50, "Status": 13,
        "Created": 17, "Updated": 17, "Lead Source": 20,
    }
    for index, col in enumerate(CRM_COLUMNS, start=1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = widths.get(col, 18)
        for cell in ws[letter][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col in ("Primary Phone", "Other Phones"):
                cell.number_format = "@"          # keep phone numbers as text

    ws.freeze_panes = "C2"                        # header + company stay visible
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"vendors_{datetime.now():%Y-%m-%d}.xlsx"
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _api_vendor(v):
    keep = ALL_FIELDS + ["id", "status", "created_at", "updated_at"]
    return {k: v.get(k) for k in keep}


@app.route("/api/vendors")
def api_vendors():
    """JSON list for CRM integrations. Optional ?search= and ?status=."""
    vendors = get_vendors(request.args.get("search", "").strip())
    status = request.args.get("status")
    if status:
        vendors = [v for v in vendors if v["status"] == status]
    return jsonify([_api_vendor(v) for v in vendors])


@app.route("/api/vendors/<int:vendor_id>")
def api_vendor(vendor_id):
    vendor = get_vendor(vendor_id)
    if not vendor:
        return jsonify({"error": "not found"}), 404
    return jsonify(_api_vendor(vendor))


# ------------------------------------------------------------- import

@app.route("/import", methods=["GET", "POST"])
def import_page():
    """Step 1: upload an Excel / CSV file."""
    if request.method == "POST":
        upload = request.files.get("file")
        if not upload or not upload.filename:
            flash("Please choose an Excel or CSV file.", "error")
            return redirect(url_for("import_page"))

        suffix = Path(upload.filename).suffix.lower()
        if suffix not in (".xlsx", ".xlsm", ".csv", ".txt", ".xls"):
            flash("Please upload an Excel (.xlsx) or CSV file.", "error")
            return redirect(url_for("import_page"))

        token = uuid.uuid4().hex
        path = IMPORT_DIR / f"{token}{suffix}"
        upload.save(path)
        try:
            read_rows(path)                      # check it can be read
        except ImportFileError as error:
            path.unlink(missing_ok=True)
            flash(str(error), "error")
            return redirect(url_for("import_page"))

        (IMPORT_DIR / f"{token}.name").write_text(upload.filename, encoding="utf-8")
        return redirect(url_for("import_preview", token=token))

    return render_template("import.html")


def _import_file(token):
    if not token.isalnum():
        abort(404)
    for path in IMPORT_DIR.glob(f"{token}.*"):
        if path.suffix != ".name":
            return path
    return None


def _original_name(token):
    name_file = IMPORT_DIR / f"{token}.name"
    return name_file.read_text(encoding="utf-8") if name_file.exists() else "file"


def _vendor_label(vendor):
    parts = [vendor.get("company") or "", vendor.get("person_name") or ""]
    return " - ".join(p for p in parts if p) or "(no name)"


def _db_matcher():
    """Matcher loaded with every saved vendor."""
    matcher = Matcher()
    for v in get_vendors():
        matcher.add(v, ("db", v["id"]))
    return matcher


def plan_import(headers, rows, mapping):
    """Dry run: what would happen to every row (nothing is saved)."""
    matcher = _db_matcher()
    plans = []
    for number, row in enumerate(rows, start=1):
        vendor, problem = row_to_vendor(headers, row, mapping)
        plan = {"number": number, "vendor": vendor, "problem": problem,
                "match": None, "reason": None}
        if not problem:
            entry, reason = matcher.find(vendor)
            if entry:
                kind, ref = entry["ref"]
                plan["match"] = {
                    "kind": kind, "ref": ref,
                    "label": _vendor_label(entry["vendor"]),
                    "vendor": dict(entry["vendor"]),
                }
                plan["reason"] = reason
                merged = merge_records(entry["vendor"], vendor)
                adds = []
                for field, word in (("phone_numbers", "phone"), ("emails", "email"),
                                    ("destinations_covered", "destination"),
                                    ("websites", "website")):
                    extra = len(merged[field]) - len(entry["vendor"].get(field) or [])
                    if extra > 0:
                        adds.append(f"{extra} {word}{'s' if extra > 1 else ''}")
                plan["adds"] = adds
                matcher.update(entry, merged)
            else:
                matcher.add(vendor, ("row", number))
        plans.append(plan)
    return plans


def run_import(headers, rows, mapping, actions, merge_mode, mark_reviewed):
    """Save the rows. actions = {row number: 'merge' | 'new' | 'skip'}."""
    matcher = _db_matcher()
    result = {"added": 0, "merged": 0, "skipped_dup": 0, "invalid": []}

    for number, row in enumerate(rows, start=1):
        vendor, problem = row_to_vendor(headers, row, mapping)
        if problem:
            result["invalid"].append((number, problem))
            continue

        entry, _ = matcher.find(vendor)
        action = actions.get(number, "merge") if entry else "new"

        if action == "skip":
            result["skipped_dup"] += 1
            continue

        if action == "merge":
            vendor_id = entry["ref"][1]
            existing = get_vendor(vendor_id)
            merged = merge_records(existing, vendor, merge_mode)
            if mark_reviewed:
                reviewed = 1
            elif merge_mode == "fill":
                reviewed = existing.get("reviewed") or 0
            else:
                reviewed = 0
            with sqlite3.connect(DB_PATH) as conn:
                update_vendor(conn, vendor_id, merged, existing.get("ocr_text") or "",
                              existing.get("ai_status") or "import")
                conn.execute("UPDATE vendors SET reviewed = ? WHERE id = ?",
                             (reviewed, vendor_id))
            matcher.update(entry, merged)
            result["merged"] += 1
            continue

        with sqlite3.connect(DB_PATH) as conn:
            cursor = conn.execute(
                f"INSERT INTO vendors ({', '.join(ALL_FIELDS)}, ocr_text, ai_status, "
                "reviewed, source, created_by) "
                f"VALUES ({', '.join('?' for _ in ALL_FIELDS)}, '', 'import', ?, 'import', ?)",
                _field_values(vendor) + [1 if mark_reviewed else 0, current_user()],
            )
            new_id = cursor.lastrowid
        matcher.add(vendor, ("db", new_id))
        result["added"] += 1

    return result


@app.route("/import/<token>", methods=["GET", "POST"])
def import_preview(token):
    """Step 2: check columns and duplicates, then import."""
    path = _import_file(token)
    if not path:
        flash("The uploaded file has expired. Please upload it again.", "error")
        return redirect(url_for("import_page"))

    try:
        headers, rows = read_rows(path)
    except ImportFileError as error:
        flash(str(error), "error")
        return redirect(url_for("import_page"))

    if request.method == "POST":
        mapping = {i: request.form.get(f"col_{i}", "") for i in range(len(headers))}
    else:
        mapping = {i: guess_target(h) for i, h in enumerate(headers)}

    if request.method == "POST" and request.form.get("action") == "import":
        if not any(t in ("company", "person_name", "first_name", "emails", "phone_numbers")
                   for t in mapping.values()):
            flash("Choose at least one column for Company, Contact name, Email or Phone.", "error")
        else:
            actions = {}
            for key, value in request.form.items():
                if key.startswith("dup_") and value in ("merge", "new", "skip"):
                    actions[int(key[4:])] = value
            merge_mode = "fill" if request.form.get("merge_mode") == "fill" else "update"
            result = run_import(headers, rows, mapping, actions, merge_mode,
                                request.form.get("mark_reviewed") == "1")

            filename = _original_name(token)
            for leftover in IMPORT_DIR.glob(f"{token}.*"):
                leftover.unlink(missing_ok=True)
            return render_template("import_done.html", filename=filename,
                                   total=len(rows), merge_mode=merge_mode, **result)

    plans = plan_import(headers, rows, mapping)
    duplicates = [p for p in plans if p["match"]]
    counts = {
        "new": sum(1 for p in plans if not p["problem"] and not p["match"]),
        "dup_saved": sum(1 for p in duplicates if p["match"]["kind"] == "db"),
        "dup_file": sum(1 for p in duplicates if p["match"]["kind"] == "row"),
        "invalid": sum(1 for p in plans if p["problem"]),
    }
    return render_template(
        "import_preview.html", token=token, filename=_original_name(token),
        headers=headers, rows=rows, mapping=[mapping[i] for i in range(len(headers))],
        targets=TARGETS,
        samples=[[r[i] for r in rows[:3]] for i in range(len(headers))],
        new_rows=[p for p in plans if not p["problem"] and not p["match"]][:8],
        duplicates=duplicates, invalid=[p for p in plans if p["problem"]],
        counts=counts, form=request.form if request.method == "POST" else {},
    )


@app.route("/import/template.xlsx")
def import_template():
    """Empty Excel file with the recommended columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    columns = ["Company Name", "Contact Name", "Job Title", "Phone", "Email",
               "Website", "Address", "Destinations", "Notes"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.append(columns)
    ws.append(["Example DMC", "Asha Mehta", "Sales Manager", "+91 98765 43210",
               "sales@example.com", "www.example.com", "Mumbai, India",
               "Dubai; Bali; Thailand", "MICE, FIT"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="09488E")
    for letter, width in zip("ABCDEFGHI", (28, 22, 20, 18, 28, 24, 30, 28, 28)):
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            cell.number_format = "@"
    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=vendor_import_template.xlsx"},
    )


# ------------------------------------------------------------ settings

@app.route("/settings", methods=["GET", "POST"])
def settings_page():
    if request.method == "POST":
        quality = request.form.get("scan_quality", "fast")
        save_settings({"scan_quality": quality if quality in SCAN_QUALITY else "fast"})
        flash("Settings saved.", "success")
        return redirect(url_for("settings_page"))

    import os
    key = os.getenv("GEMINI_API_KEY") or ""
    with sqlite3.connect(DB_PATH) as conn:
        record_count = conn.execute("SELECT COUNT(*) FROM vendors").fetchone()[0]
    images = [p for p in UPLOAD_DIR.glob("*.jpg")]
    info = {
        "api_key": ("•••• " + key[-4:]) if len(key) > 8 else "",
        "db_path": str(DB_PATH),
        "db_size": DB_PATH.stat().st_size if DB_PATH.exists() else 0,
        "upload_path": str(UPLOAD_DIR),
        "image_count": len(images),
        "image_size": sum(p.stat().st_size for p in images),
        "record_count": record_count,
        "fallback_models": ai_interpreter.MODELS[1:],
    }
    return render_template("settings.html", settings=get_settings(), info=info)


@app.route("/settings/backup")
def download_backup():
    """Download a copy of the database (safe to do while the app runs)."""
    with sqlite3.connect(DB_PATH) as conn:
        data = conn.serialize()
    filename = f"vendors_backup_{datetime.now():%Y-%m-%d_%H%M}.db"
    return Response(
        data, mimetype="application/octet-stream",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.template_filter("filesize")
def filesize(num):
    for unit in ("bytes", "KB", "MB", "GB"):
        if num < 1024 or unit == "GB":
            return f"{num:.0f} {unit}" if unit == "bytes" else f"{num:.1f} {unit}"
        num /= 1024


@app.context_processor
def layout_data():
    """Values every page needs for the sidebar."""
    settings = get_settings()
    with sqlite3.connect(DB_PATH) as conn:
        total = conn.execute("SELECT COUNT(*) FROM vendors").fetchone()[0]
        to_review = conn.execute(
            "SELECT COUNT(*) FROM vendors WHERE COALESCE(reviewed, 0) = 0"
        ).fetchone()[0]
    return {
        "app_settings": settings,
        "nav_counts": {"total": total, "to_review": to_review},
    }


# ------------------------------------------------- Excel (Microsoft 365) sync

TOKEN_CACHE = BASE_DIR / "ms_token_cache.json"
_sync_lock = threading.Lock()
_auth_cache = {}


def ms_auth():
    """One sign-in helper per Client/Tenant ID (keeps the sign-in progress)."""
    settings = get_settings()
    key = (settings["ms_client_id"], settings["ms_tenant_id"])
    if key not in _auth_cache:
        _auth_cache.clear()
        _auth_cache[key] = sheet_sync.MicrosoftAuth(key[0], key[1], TOKEN_CACHE)
    return _auth_cache[key]


class VendorStore:
    """Gives sheet_sync access to the vendors in the database."""

    def vendors(self):
        return get_vendors()

    def insert(self, vendor):
        with sqlite3.connect(DB_PATH) as conn:
            cursor = conn.execute(
                f"INSERT INTO vendors ({', '.join(ALL_FIELDS)}, ocr_text, ai_status, "
                "reviewed, source) "
                f"VALUES ({', '.join('?' for _ in ALL_FIELDS)}, '', 'import', 0, 'sheet')",
                _field_values(vendor),
            )
            return cursor.lastrowid

    def update(self, vendor_id, vendor):
        existing = get_vendor(vendor_id)
        if not existing:
            return
        merged = {f: vendor.get(f, existing.get(f)) for f in ALL_FIELDS}
        with sqlite3.connect(DB_PATH) as conn:
            update_vendor(conn, vendor_id, merged, existing.get("ocr_text") or "",
                          existing.get("ai_status") or "ok")

    def links(self):
        with sqlite3.connect(DB_PATH) as conn:
            conn.row_factory = sqlite3.Row
            return {row["vendor_id"]: dict(row) for row in conn.execute("SELECT * FROM sheet_links")}

    def set_link(self, vendor_id, app_hash, sheet_hash):
        with sqlite3.connect(DB_PATH) as conn:
            conn.execute(
                "INSERT INTO sheet_links (vendor_id, app_hash, sheet_hash, deleted, removed, synced_at) "
                "VALUES (?, ?, ?, 0, 0, CURRENT_TIMESTAMP) "
                "ON CONFLICT(vendor_id) DO UPDATE SET app_hash = excluded.app_hash, "
                "sheet_hash = excluded.sheet_hash, deleted = 0, removed = 0, "
                "synced_at = CURRENT_TIMESTAMP",
                (vendor_id, app_hash, sheet_hash),
            )

    def mark_removed(self, vendor_id):
        with sqlite3.connect(DB_PATH) as conn:
            conn.execute("UPDATE sheet_links SET removed = 1 WHERE vendor_id = ?", (vendor_id,))


def run_sheet_sync(trigger="manual"):
    """Run one sync and store the result. Returns (ok, message)."""
    if not _sync_lock.acquire(blocking=False):
        return False, "A sync is already running - please wait a moment."
    started = datetime.now()
    try:
        settings = get_settings()
        try:
            cached = json.loads(settings["sheet_ids"] or "{}")
        except ValueError:
            cached = {}
        sheet = sheet_sync.ExcelOnlineSheet(
            ms_auth(), settings["sheet_url"], settings["sheet_worksheet"], cached,
            save_ids=lambda ids: save_settings({"sheet_ids": json.dumps(ids)}),
        )
        report = sheet_sync.run_sync(sheet, VendorStore(), log=lambda t: print("[sync]", t))
        report["trigger"] = trigger
        report["seconds"] = round((datetime.now() - started).total_seconds())
        changes = (report["to_sheet_added"] + report["to_sheet_updated"] + report["to_app_added"]
                   + report["to_app_updated"] + report["to_app_merged"])
        message = ("Sync finished - everything was already up to date." if not changes else
                   f"Sync finished: {report['to_sheet_added']} added and {report['to_sheet_updated']} "
                   f"updated in the sheet; {report['to_app_added']} added and "
                   f"{report['to_app_updated'] + report['to_app_merged']} updated in the app.")
        ok = True
    except sheet_sync.SyncError as error:
        report, message, ok = {"error": str(error), "trigger": trigger}, f"Sync failed: {error}", False
    except Exception as error:                      # network down etc.
        app.logger.exception("Excel sync failed")
        report, message, ok = {"error": str(error), "trigger": trigger}, f"Sync failed: {error}", False
    finally:
        _sync_lock.release()

    save_settings({
        "sync_last_at": started.strftime("%Y-%m-%d %H:%M"),
        "sync_last_ok": "1" if ok else "0",
        "sync_last_report": json.dumps({**report, "message": message}),
    })
    print("[sync]", message)
    return ok, message


def sync_scheduler():
    """Background thread: run the daily sync once a day after the set hour
    (or as soon as the app is started after that hour)."""
    while True:
        try:
            settings = get_settings()
            now = datetime.now()
            today = now.strftime("%Y-%m-%d")
            hour = int(settings["sync_hour"] or 10)
            ready = (settings["sync_enabled"] == "1" and settings["ms_client_id"]
                     and TOKEN_CACHE.exists())
            if ready and now.hour >= hour and settings["sync_last_auto_date"] != today:
                save_settings({"sync_last_auto_date": today})
                run_sheet_sync(trigger="daily")
        except Exception:
            app.logger.exception("Sync scheduler error")
        time.sleep(60)


@app.route("/sync", methods=["GET", "POST"])
def sync_page():
    if request.method == "POST":
        hour = request.form.get("sync_hour", "10")
        save_settings({
            "ms_client_id": request.form.get("ms_client_id", "").strip(),
            "ms_tenant_id": request.form.get("ms_tenant_id", "").strip(),
            "sheet_url": request.form.get("sheet_url", "").strip(),
            "sheet_worksheet": request.form.get("sheet_worksheet", "").strip() or "CARD_DETAILS",
            "sync_enabled": "1" if request.form.get("sync_enabled") else "0",
            "sync_hour": hour if hour.isdigit() and 0 <= int(hour) <= 23 else "10",
        })
        flash("Excel sync settings saved.", "success")
        return redirect(url_for("sync_page"))

    settings = get_settings()
    auth = ms_auth()
    try:
        last = json.loads(settings["sync_last_report"] or "{}")
    except ValueError:
        last = {}
    try:
        import msal  # noqa: F401
        msal_ok = True
    except ImportError:
        msal_ok = False
    account = auth.account() if TOKEN_CACHE.exists() else None
    hour = int(settings["sync_hour"] or 10)
    now = datetime.now()
    next_run = now.replace(hour=hour, minute=0, second=0, microsecond=0)
    if settings["sync_last_auto_date"] == now.strftime("%Y-%m-%d") or now >= next_run:
        next_run = next_run + timedelta(days=1)
    with sqlite3.connect(DB_PATH) as conn:
        linked = conn.execute("SELECT COUNT(*) FROM sheet_links WHERE deleted = 0").fetchone()[0]
    return render_template(
        "sync.html", settings=settings, account=account, last=last, msal_ok=msal_ok,
        configured=bool(settings["ms_client_id"] and settings["sheet_url"]),
        next_run=next_run.strftime("%d %b %Y, %H:%M"), linked=linked,
        flow=auth.flow_state,
    )


@app.route("/sync/signin", methods=["POST"])
def sync_signin():
    try:
        ms_auth().start_sign_in()
    except sheet_sync.SyncError as error:
        flash(str(error), "error")
    return redirect(url_for("sync_page"))


@app.route("/sync/signin/status")
def sync_signin_status():
    state = dict(ms_auth().flow_state)
    if state.get("state") == "done":
        state["account"] = ms_auth().account()
    return jsonify(state)


@app.route("/sync/signout", methods=["POST"])
def sync_signout():
    ms_auth().sign_out()
    flash("Signed out of Microsoft 365. The daily sync is paused until you sign in again.", "info")
    return redirect(url_for("sync_page"))


@app.route("/sync/run", methods=["POST"])
def sync_run():
    ok, message = run_sheet_sync(trigger="manual")
    flash(message, "success" if ok else "error")
    return redirect(url_for("sync_page"))


@app.route("/retry/<int:vendor_id>", methods=["POST"])
def retry_ai(vendor_id):
    """Re-run Gemini on the saved OCR text (no new scan needed)."""
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT * FROM vendors WHERE id = ?", (vendor_id,)).fetchone()

    if not row or not row["ocr_text"]:
        flash("No saved card text for this record - please scan the card again.", "error")
        return redirect(url_for("index"))

    try:
        start = time.perf_counter()
        vendor = interpret_vendor_text(row["ocr_text"])
        with sqlite3.connect(DB_PATH) as conn:
            update_vendor(conn, vendor_id, vendor, row["ocr_text"], "ok")
        name = vendor.get("company") or vendor.get("person_name") or "Vendor"
        flash(f"{name}: AI details added ({time.perf_counter() - start:.0f}s). "
              "Please review them.", "success")
    except AIUnavailable:
        flash("The AI service is still busy. Please try again in a few minutes.", "warning")
    except Exception as error:
        app.logger.exception("AI retry failed")
        flash(f"AI retry failed: {error}", "error")

    return redirect(url_for("vendor_detail", vendor_id=vendor_id))


@app.errorhandler(413)
def too_large(_error):
    flash("That file is too large (max 25 MB).", "error")
    return redirect(request.referrer or url_for("index"))


def warm_up_ocr():
    """The first OCR call is slow; do it at startup instead of on the
    first card."""
    import numpy as np
    start = time.perf_counter()
    ocr.predict(np.full((64, 256, 3), 255, dtype=np.uint8))
    print(f"OCR ready ({time.perf_counter() - start:.1f}s warm-up)")


init_db()
apply_settings()


def lan_address():
    """This PC's address on the office network (e.g. 192.168.1.25)."""
    import socket
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sock:
            sock.connect(("8.8.8.8", 80))       # no data is sent
            return sock.getsockname()[0]
    except OSError:
        return socket.gethostbyname(socket.gethostname())


if __name__ == "__main__":
    import os
    import sys

    # Run "python app.py --local" to allow only this PC.
    local_only = "--local" in sys.argv
    host = "127.0.0.1" if local_only else "0.0.0.0"
    port = int(os.getenv("PORT", "5000"))

    warm_up_ocr()
    threading.Thread(target=sync_scheduler, daemon=True).start()
    print("\n" + "=" * 60)
    print(f"  On this PC:          http://127.0.0.1:{port}")
    if not local_only:
        print(f"  Other office PCs:    http://{lan_address()}:{port}")
    print("=" * 60 + "\n")

    # debug=False: the debug mode lets anyone who can open the page run
    # commands on this PC, so it must be off when others can connect.
    # use_reloader=False: avoids loading the OCR models twice.
    app.run(host=host, port=port, debug=False, use_reloader=False, threaded=True)
