"""
Excel / CSV import.

Reads a spreadsheet, guesses which column holds what (company, name,
phone, email ...) from the header names, and turns every row into a
vendor record. The user can correct the guessed columns on the preview
page before anything is saved.
"""

import csv
import io
import re
from pathlib import Path

# Fields a column can be mapped to (value -> label shown in the dropdown)
TARGETS = {
    "": "— Ignore this column —",
    "company": "Company",
    "person_name": "Contact name (full)",
    "first_name": "First name",
    "last_name": "Last name",
    "designation": "Job title",
    "phone_numbers": "Phone",
    "emails": "Email",
    "websites": "Website",
    "address": "Address / city / country",
    "destinations_covered": "Destinations",
    "services_notes": "Notes / services",
    "other_contacts": "Other contacts",
}

# Header names we recognise (compared lower-case, letters and digits only)
SYNONYMS = {
    "company": [
        "company", "companyname", "organisation", "organization", "org",
        "vendor", "vendorname", "firm", "agency", "agencyname", "business",
        "businessname", "account", "accountname", "supplier", "suppliername",
        "dmc", "dmcname", "partner", "partnername",
    ],
    "person_name": [
        "name", "fullname", "contactname", "contact", "contactperson",
        "person", "personname", "keycontact", "representative",
    ],
    "first_name": ["firstname", "first", "givenname", "fname"],
    "last_name": ["lastname", "surname", "last", "familyname", "lname"],
    "designation": [
        "designation", "title", "jobtitle", "position", "role", "post",
    ],
    "phone_numbers": [
        "phone", "phones", "mobile", "mobileno", "mobilenumber", "phoneno",
        "phonenumber", "contactno", "contactnumber", "tel", "telephone",
        "whatsapp", "whatsappno", "primaryphone", "otherphones", "landline",
        "cell", "cellphone", "officephone", "directline", "mob",
    ],
    "emails": [
        "email", "emails", "emailid", "mail", "emailaddress", "primaryemail",
        "otheremails", "mailid", "officialemail",
    ],
    "websites": ["website", "websites", "web", "url", "site", "webaddress"],
    "address": [
        "address", "officeaddress", "city", "state", "country", "location",
        "pincode", "pin", "zip", "zipcode", "postcode", "postalcode",
        "street", "addressline1", "addressline2", "address1", "address2",
        "region",
    ],
    "destinations_covered": [
        "destinations", "destination", "destinationscovered", "countries",
        "countriescovered", "coverage", "markets", "regionscovered",
        "operatingcountries",
    ],
    "services_notes": [
        "notes", "note", "remarks", "remark", "services", "service",
        "comments", "comment", "description", "servicesnotes", "products",
        "speciality", "specialty",
    ],
    "other_contacts": ["othercontacts", "additionalcontacts"],
    # Columns from our own export that should not be imported as data
    "": [
        "recordid", "id", "status", "created", "updated", "leadsource",
        "sno", "srno", "serialno", "slno", "no", "sr",
    ],
}

_LOOKUP = {}
for _target, _names in SYNONYMS.items():
    for _name in _names:
        _LOOKUP.setdefault(_name, _target)

_EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")


def _key(header):
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def guess_target(header):
    """Best field for a column header, or '' to ignore."""
    key = _key(header)
    if not key:
        return ""
    if key in _LOOKUP:
        return _LOOKUP[key]
    # Strong words win, e.g. "Contact Email" is an email, not a name
    for words, target in (
        (("email", "mail"), "emails"),
        (("phone", "mobile", "whatsapp", "telephone", "contactno", "contactnumber"), "phone_numbers"),
        (("website", "url"), "websites"),
        (("destination",), "destinations_covered"),
    ):
        if any(w in key for w in words):
            return target
    # Partial matches, e.g. "Vendor Company Name", "Mobile No (WhatsApp)"
    best, best_len = "", 0
    for name, target in _LOOKUP.items():
        if len(name) >= 4 and name in key and len(name) > best_len:
            best, best_len = target, len(name)
    return best


def _cell_text(value):
    """Excel cell -> clean text (numbers keep no '.0', dates as text)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


class ImportFileError(ValueError):
    """The file could not be read."""


def read_rows(path):
    """Return (headers, rows) from .xlsx / .xlsm / .csv / .txt."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:
            raise ImportFileError(f"This Excel file could not be opened ({error}).")
        # Use the sheet with the most filled rows
        best = None
        for ws in wb.worksheets:
            data = [
                [_cell_text(c) for c in row]
                for row in ws.iter_rows(values_only=True)
            ]
            data = [r for r in data if any(r)]
            if best is None or len(data) > len(best):
                best = data
        wb.close()
        table = best or []

    elif suffix in (".csv", ".txt"):
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        table = [
            [c.strip() for c in row]
            for row in csv.reader(io.StringIO(text), dialect)
        ]
        table = [r for r in table if any(r)]

    elif suffix == ".xls":
        raise ImportFileError(
            "Old .xls files are not supported. In Excel choose "
            "File > Save As > Excel Workbook (.xlsx) and upload that file."
        )
    else:
        raise ImportFileError("Please upload an Excel (.xlsx) or CSV file.")

    if not table:
        raise ImportFileError("The file is empty.")

    # Header row = the first of the top 10 rows that looks most like headers
    header_index, best_score = 0, -1
    for i, row in enumerate(table[:10]):
        score = sum(1 for c in row if c and guess_target(c))
        if score > best_score:
            header_index, best_score = i, score

    headers = table[header_index]
    width = max(len(r) for r in table)
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers + [""] * (width - len(headers)))]
    rows = [r + [""] * (width - len(r)) for r in table[header_index + 1:]]
    return headers, rows


def _split(value, pattern):
    return [p.strip() for p in re.split(pattern, value) if p.strip()]


def row_to_vendor(headers, row, mapping):
    """Turn one spreadsheet row into a vendor dict using the column mapping
    {column index: target field}. Returns (vendor, problem) - problem is a
    text reason when the row should be skipped."""
    vendor = {
        "company": "", "person_name": "", "designation": "",
        "phone_numbers": [], "emails": [], "websites": [], "address": "",
        "destinations_covered": [], "services_notes": "", "other_contacts": [],
    }
    first, last, address_parts, notes = "", "", [], []

    for index, target in mapping.items():
        if not target or index >= len(row):
            continue
        value = " ".join(str(row[index]).split()) if target not in (
            "phone_numbers", "emails", "destinations_covered", "other_contacts", "websites"
        ) else str(row[index]).strip()
        if not value:
            continue

        if target == "first_name":
            first = value
        elif target == "last_name":
            last = value
        elif target in ("company", "person_name", "designation"):
            vendor[target] = vendor[target] or value
        elif target == "address":
            address_parts.append(value)
        elif target == "services_notes":
            notes.append(value)
        elif target == "phone_numbers":
            for phone in _split(value, r"[\n;,|]+"):
                phone = re.sub(r"^\s*(?:[A-Za-z]{1,3}\s*[:.]\s*)", "", phone).strip()
                if len(re.sub(r"\D", "", phone)) >= 6 and phone not in vendor["phone_numbers"]:
                    vendor["phone_numbers"].append(phone)
        elif target == "emails":
            for email in _split(value, r"[\s;,|]+"):
                email = email.lower().strip(".")
                if _EMAIL_RE.match(email):
                    if email not in vendor["emails"]:
                        vendor["emails"].append(email)
                else:
                    notes.append(f"Unrecognised email: {email}")
        elif target == "websites":
            for site in _split(value, r"[\s;,|]+"):
                site = site.lower()
                if "." in site and site not in vendor["websites"]:
                    vendor["websites"].append(site)
        elif target == "destinations_covered":
            for place in _split(value, r"[\n;,|•/]+"):
                if place not in vendor["destinations_covered"]:
                    vendor["destinations_covered"].append(place)
        elif target == "other_contacts":
            for item in _split(value, r"[\n|]+"):
                vendor["other_contacts"].append(item)

    if not vendor["person_name"] and (first or last):
        vendor["person_name"] = " ".join(p for p in (first, last) if p)

    seen = []
    for part in address_parts:
        if part.lower() not in " ".join(seen).lower():
            seen.append(part)
    vendor["address"] = ", ".join(seen)
    vendor["services_notes"] = "; ".join(notes)

    if not (vendor["company"] or vendor["person_name"]):
        if vendor["emails"] or vendor["phone_numbers"]:
            # Keep contacts without a name, use the email domain as company
            if vendor["emails"]:
                vendor["company"] = vendor["emails"][0].split("@")[1].split(".")[0].upper()
        else:
            return vendor, "no company, name, phone or email"

    return vendor, None
